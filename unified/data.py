from __future__ import annotations

import csv
import json
import math
import random
import re
import uuid
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path

import blosc2
import numpy as np
import torch
import torch.nn.functional as F

from unified.specs import TaskSpec, canonical_modality_id, discover_task_specs, safe_key


def scaled_patch_size(planned: list[int], target_voxels: int, divisor: int = 16) -> tuple[int, int, int]:
    scale = min(1.0, (target_voxels / math.prod(planned)) ** (1.0 / 3.0))
    return tuple(max(divisor, int(value * scale) // divisor * divisor) for value in planned)


@dataclass
class PreparedTask:
    spec: TaskSpec
    raw_dir: Path
    data_dir: Path
    identifiers: list[str]
    modality_ids: tuple[int, ...]
    labels: dict[str, dict[str, int]]
    clinical: dict[str, np.ndarray]
    patch_size: tuple[int, int, int]


def _read_classification(raw_dir: Path, spec: TaskSpec) -> dict[str, dict[str, int]]:
    with (raw_dir / "cls_data.csv").open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    labels: dict[str, dict[str, int]] = {}
    if not rows:
        return {}
    identifier_column = (
        "identifier" if "identifier" in rows[0] else "case_id" if "case_id" in rows[0] else next(iter(rows[0]))
    )
    dataset = json.loads((raw_dir / "dataset.json").read_text())
    declared_heads = list(dataset.get("classification_labels", {}))
    for row in rows:
        if spec.key == "glioma_idh_grade":
            labels[row["identifier"]] = {"idh": int(row["idh"]), "grade": int(row["grade"])}
        else:
            identifier = row.get(identifier_column)
            if not identifier:
                raise ValueError(f"{spec.key}: empty identifier in cls_data.csv")
            values = {}
            for index, head_name in enumerate(spec.classification_heads):
                declared = declared_heads[index] if index < len(declared_heads) else head_name
                candidates = (head_name, declared, safe_key(declared), "label" if len(spec.classification_heads) == 1 else "")
                source_column = next((candidate for candidate in candidates if candidate and candidate in row), None)
                if source_column is None:
                    raise ValueError(
                        f"{spec.key}: cannot map classification head {head_name!r} to columns {list(row)}"
                    )
                values[head_name] = int(row[source_column])
            labels[identifier] = values
    return labels


def _read_clinical(
    raw_dir: Path,
    identifiers: list[str],
    dimension: int = 8,
    fit_identifiers: list[str] | None = None,
) -> dict[str, np.ndarray]:
    """Automatically encode numeric/categorical clinical columns into a fixed vector."""
    zeros = {identifier: np.zeros(dimension, dtype=np.float32) for identifier in identifiers}
    csv_path = raw_dir / "clinical_data.csv"
    xlsx_path = raw_dir / "clinical_data.xlsx"
    if csv_path.exists():
        with csv_path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
    elif xlsx_path.exists():
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError("Reading clinical_data.xlsx requires openpyxl") from error
        workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(values)]
        rows = [
            {header: "" if value is None else str(value).strip() for header, value in zip(headers, row)}
            for row in values
        ]
        workbook.close()
    else:
        return zeros
    if not rows:
        return zeros

    def normalize(value: object) -> str:
        return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())

    aliases: dict[str, dict[str, str] | None] = {}
    for row in rows:
        values = [value for value in row.values() if str(value).strip()]
        candidates = {normalize(value) for value in values}
        candidates |= {"".join(re.findall(r"\d+", str(value))) for value in values}
        # Composite identifiers such as patient_id + study_id are common.
        for left in range(len(values)):
            for right in range(left + 1, len(values)):
                candidates.add(normalize(values[left]) + normalize(values[right]))
                candidates.add("".join(re.findall(r"\d+", str(values[left]) + str(values[right]))))
        for alias in (value for value in candidates if value):
            if alias in aliases and aliases[alias] is not row:
                aliases[alias] = None
            else:
                aliases[alias] = row

    def match_row(identifier: str) -> dict[str, str]:
        normalized = normalize(identifier)
        numeric = "".join(re.findall(r"\d+", identifier))
        for candidate in (normalized, numeric, re.sub(r"^[a-z]+", "", normalized)):
            match = aliases.get(candidate)
            if match is not None:
                return match
        return {}

    by_identifier = {identifier: match_row(identifier) for identifier in identifiers}
    excluded = re.compile(
        r"(^|_)(id|uid)($|_)|date|time|coord|label|grade|isup|diagnos|disease|histopath|lesion|cancer|malignan|annotation|series",
        flags=re.IGNORECASE,
    )
    feature_columns = [column for column in rows[0] if not excluded.search(column)][:dimension]
    encoded_columns: list[dict[str, float] | tuple[float, float]] = []
    fitting = identifiers if fit_identifiers is None else fit_identifiers
    training_rows = [by_identifier[value] for value in fitting if by_identifier.get(value)]
    for column in feature_columns:
        values = [row.get(column, "").strip() for row in training_rows]
        numeric_values = []
        numeric = True
        for value in values:
            try:
                numeric_values.append(float(value))
            except (TypeError, ValueError):
                numeric = False
                break
        if numeric and numeric_values:
            center = float(np.median(numeric_values))
            scale = float(np.std(numeric_values))
            encoded_columns.append((center, max(scale, 1e-6)))
        else:
            categories = sorted({value for value in values if value})
            if len(categories) <= 1:
                mapping = {value: 0.0 for value in categories}
            else:
                mapping = {value: 2.0 * index / (len(categories) - 1) - 1.0 for index, value in enumerate(categories)}
            encoded_columns.append(mapping)
    output = {}
    for identifier in identifiers:
        vector = np.zeros(dimension, dtype=np.float32)
        row = by_identifier.get(identifier, {})
        for index, (column, encoder) in enumerate(zip(feature_columns, encoded_columns)):
            value = row.get(column, "").strip()
            if isinstance(encoder, tuple):
                try:
                    vector[index] = np.clip((float(value) - encoder[0]) / encoder[1], -5, 5)
                except (TypeError, ValueError):
                    vector[index] = 0
            else:
                vector[index] = encoder.get(value, 0.0)
        output[identifier] = vector
    return output


def discover_tasks(
    raw_root: Path, preprocessed_root: Path, target_patch_voxels: int = 2_000_000
) -> dict[str, PreparedTask]:
    output: dict[str, PreparedTask] = {}
    for key, spec in discover_task_specs(raw_root).items():
        raw_dir = raw_root / spec.folder_name
        preprocessed_dir = preprocessed_root / spec.folder_name
        plans_path = preprocessed_dir / "nnUNetPlans.json"
        if not plans_path.exists():
            continue
        plans = json.loads(plans_path.read_text())
        configuration = plans["configurations"]["3d_fullres"]
        data_dir = preprocessed_dir / configuration["data_identifier"]
        identifiers = sorted(path.name.removesuffix(".b2nd") for path in data_dir.glob("*.b2nd") if not path.name.endswith("_seg.b2nd"))
        dataset_json = json.loads((raw_dir / "dataset.json").read_text())
        channels = dataset_json["channel_names"]
        modality_ids = tuple(canonical_modality_id(channels[str(index)]) for index in range(len(channels)))
        labels = _read_classification(raw_dir, spec)
        missing = set(identifiers) - set(labels)
        if missing:
            raise ValueError(f"{key}: {len(missing)} preprocessed cases have no classification label")
        output[key] = PreparedTask(
            spec=spec,
            raw_dir=raw_dir,
            data_dir=data_dir,
            identifiers=identifiers,
            modality_ids=modality_ids,
            labels=labels,
            clinical=_read_clinical(raw_dir, identifiers),
            # A modality is encoded independently before fusion, so constrain total
            # channel-voxels rather than silently assuming a single-channel input.
            patch_size=scaled_patch_size(
                configuration["patch_size"], max(16**3, target_patch_voxels // len(modality_ids))
            ),
        )
    return output


def _crop_or_pad(array: np.ndarray, center: np.ndarray, patch_size: tuple[int, int, int]) -> np.ndarray:
    spatial = np.asarray(array.shape[-3:])
    patch = np.asarray(patch_size)
    start = center - patch // 2
    end = start + patch
    source_start = np.maximum(start, 0)
    source_end = np.minimum(end, spatial)
    destination_start = source_start - start
    destination_end = destination_start + (source_end - source_start)
    output = np.zeros((*array.shape[:-3], *patch_size), dtype=array.dtype)
    source = tuple(slice(int(a), int(b)) for a, b in zip(source_start, source_end))
    destination = tuple(slice(int(a), int(b)) for a, b in zip(destination_start, destination_end))
    output[(..., *destination)] = array[(..., *source)]
    return output


def _read_blosc(path: Path) -> np.ndarray:
    array = blosc2.open(urlpath=str(path), mode="r", dparams={"nthreads": 1})
    return np.asarray(array[:])


def _resize_and_pad(array: np.ndarray, output_size: int, mode: str) -> np.ndarray:
    spatial = np.asarray(array.shape[-3:])
    scale = min(output_size / spatial)
    resized_shape = np.maximum(1, np.round(spatial * scale).astype(int))
    tensor = torch.from_numpy(np.ascontiguousarray(array))[None].float()
    if mode == "nearest":
        resized = F.interpolate(tensor, size=tuple(resized_shape), mode=mode)[0]
    else:
        resized = F.interpolate(tensor, size=tuple(resized_shape), mode=mode, align_corners=False)[0]
    padding = []
    for size in reversed(resized_shape):
        total = output_size - int(size)
        padding.extend((total // 2, total - total // 2))
    return F.pad(resized, padding).numpy()


class TaskPatchLoader:
    def __init__(
        self,
        tasks: dict[str, PreparedTask],
        workers: int = 8,
        seed: int = 20260824,
        balanced_fraction: float = 0.33,
        glance_size: int = 64,
        glance_cache: Path = Path("data/glance_cache"),
    ):
        self.tasks = tasks
        self.rng = random.Random(seed)
        self.balanced_fraction = balanced_fraction
        self.executor = ThreadPoolExecutor(max_workers=workers)
        self.glance_size = glance_size
        self.glance_cache = glance_cache
        self.order: dict[str, list[str]] = {}
        self.cursor: dict[str, int] = {}
        self.label_buckets: dict[str, list[list[str]]] = {}
        for key, task in tasks.items():
            self.order[key] = list(task.identifiers)
            self.rng.shuffle(self.order[key])
            self.cursor[key] = 0
            buckets: dict[tuple[int, ...], list[str]] = {}
            head_names = tuple(task.spec.classification_heads)
            for identifier in task.identifiers:
                label_key = tuple(task.labels[identifier][head] for head in head_names)
                buckets.setdefault(label_key, []).append(identifier)
            self.label_buckets[key] = list(buckets.values())

    def next_identifiers(self, task_key: str, batch_size: int) -> list[str]:
        values: list[str] = []
        balanced_count = round(batch_size * self.balanced_fraction)
        for _ in range(balanced_count):
            bucket = self.rng.choice(self.label_buckets[task_key])
            candidates = [identifier for identifier in bucket if identifier not in values]
            if not candidates:
                candidates = [identifier for identifier in self.tasks[task_key].identifiers if identifier not in values]
            values.append(self.rng.choice(candidates))
        while len(values) < batch_size:
            order = self.order[task_key]
            cursor = self.cursor[task_key]
            candidate = order[cursor]
            cursor += 1
            if cursor == len(order):
                self.rng.shuffle(order)
                cursor = 0
            self.cursor[task_key] = cursor
            if candidate not in values:
                values.append(candidate)
        return values

    def _glance_view(
        self, task_key: str, identifier: str, data: np.ndarray, segmentation: np.ndarray
    ) -> tuple[np.ndarray, np.ndarray]:
        cache_dir = self.glance_cache / task_key
        cache_dir.mkdir(parents=True, exist_ok=True)
        image_path = cache_dir / f"{identifier}_image.npy"
        label_path = cache_dir / f"{identifier}_seg.npy"
        if image_path.exists() and label_path.exists():
            return np.load(image_path).astype(np.float32), np.load(label_path).astype(np.int64)
        glance = _resize_and_pad(data, self.glance_size, "trilinear").astype(np.float16)
        glance_seg = _resize_and_pad(segmentation.astype(np.float32), self.glance_size, "nearest").astype(np.uint8)
        token = uuid.uuid4().hex
        image_tmp = image_path.with_name(f"{image_path.stem}.{token}.tmp.npy")
        label_tmp = label_path.with_name(f"{label_path.stem}.{token}.tmp.npy")
        np.save(image_tmp, glance)
        np.save(label_tmp, glance_seg)
        image_tmp.replace(image_path)
        label_tmp.replace(label_path)
        return glance.astype(np.float32), glance_seg.astype(np.int64)

    def _load_one(
        self, task_key: str, identifier: str, augment: bool, sample_seed: int
    ) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
        rng = random.Random(sample_seed)
        task = self.tasks[task_key]
        data = _read_blosc(task.data_dir / f"{identifier}.b2nd").astype(np.float32, copy=False)
        segmentation = _read_blosc(task.data_dir / f"{identifier}_seg.b2nd").astype(np.int64, copy=False)
        if segmentation.ndim == 3:
            segmentation = segmentation[None]
        glance, glance_segmentation = self._glance_view(task_key, identifier, data, segmentation)
        foreground = np.argwhere(segmentation[0] > 0)
        if len(foreground) and rng.random() < 0.67:
            center = foreground[rng.randrange(len(foreground))]
        else:
            center = np.asarray([rng.randrange(size) for size in segmentation.shape[-3:]])
        data = _crop_or_pad(data, center, task.patch_size)
        segmentation = _crop_or_pad(segmentation, center, task.patch_size)
        if augment:
            for axis in range(3):
                if rng.random() < 0.5:
                    data = np.flip(data, axis=axis + 1)
                    segmentation = np.flip(segmentation, axis=axis + 1)
            if rng.random() < 0.2:
                data = data * rng.uniform(0.9, 1.1) + rng.uniform(-0.1, 0.1)
        return (
            np.ascontiguousarray(data),
            np.ascontiguousarray(segmentation[0]),
            np.ascontiguousarray(glance),
            np.ascontiguousarray(glance_segmentation[0]),
        )

    def batch(self, task_key: str, batch_size: int, augment: bool = True) -> dict[str, object]:
        task = self.tasks[task_key]
        identifiers = self.next_identifiers(task_key, batch_size)
        seeds = [self.rng.randrange(2**63) for _ in identifiers]
        futures = [
            self.executor.submit(self._load_one, task_key, identifier, augment, seed)
            for identifier, seed in zip(identifiers, seeds)
        ]
        loaded = [future.result() for future in futures]
        images = torch.from_numpy(np.stack([item[0] for item in loaded]))
        segmentation = torch.from_numpy(np.stack([item[1] for item in loaded]))
        glance_images = torch.from_numpy(np.stack([item[2] for item in loaded]))
        glance_segmentation = torch.from_numpy(np.stack([item[3] for item in loaded]))
        modality_ids = torch.tensor(task.modality_ids, dtype=torch.long).repeat(batch_size, 1)
        classifications = {
            head: torch.tensor([task.labels[identifier][head] for identifier in identifiers], dtype=torch.long)
            for head in task.spec.classification_heads
        }
        clinical = torch.from_numpy(np.stack([task.clinical[identifier] for identifier in identifiers]))
        return {
            "task_key": task_key,
            "identifiers": identifiers,
            "images": images,
            "segmentation": segmentation,
            "glance_images": glance_images,
            "glance_segmentation": glance_segmentation,
            "modality_ids": modality_ids,
            "classifications": classifications,
            "clinical": clinical,
        }
